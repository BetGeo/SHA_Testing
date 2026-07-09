"""Reads/writes the TimeLog sheet of the existing Excel timesheet.

Only ever touches the Date / Code / Task / Hours cells of a single row.
Every other column (Project Name, Week, Billed, ...) is left alone since
those are populated by formulas already in the workbook.
"""
import time
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import openpyxl


@dataclass
class WriteResult:
    ok: bool
    message: str
    row: int | None = None


class TimesheetWriter:
    def __init__(
        self,
        path: str,
        sheet_name: str = "TimeLog",
        header_row: int = 14,
        date_col: str = "F",
        code_col: str = "G",
        task_col: str = "I",
        hours_col: str = "J",
        staff_col: str | None = "E",
    ):
        self.path = Path(path)
        self.sheet_name = sheet_name
        self.header_row = header_row
        self.date_col = date_col
        self.code_col = code_col
        self.task_col = task_col
        self.hours_col = hours_col
        self.staff_col = staff_col

    def _find_target_row(self, ws, entry_date: date, staff_name: str | None) -> tuple[int | None, str | None]:
        """Return (row_number, error_reason). error_reason is None on success."""
        candidate_row = None
        already_used_row = None
        for row in range(self.header_row + 1, ws.max_row + 1):
            cell_value = ws[f"{self.date_col}{row}"].value
            row_date = cell_value.date() if hasattr(cell_value, "date") else cell_value
            if row_date != entry_date:
                continue
            if staff_name and self.staff_col:
                staff_val = ws[f"{self.staff_col}{row}"].value
                if staff_val and staff_name.lower() not in str(staff_val).lower():
                    continue
            code_val = ws[f"{self.code_col}{row}"].value
            if code_val in (None, ""):
                candidate_row = row
                break
            already_used_row = row

        if candidate_row:
            return candidate_row, None
        if already_used_row:
            return None, (
                f"Every row for {entry_date} already has an entry (last one on row "
                f"{already_used_row}). Use the sheet's 'Insert New Row' button for a "
                f"second entry on this date, then re-run."
            )
        return None, f"No row found for {entry_date} in sheet '{self.sheet_name}'."

    def write_entry(
        self,
        entry_date: date,
        code: str,
        task: str,
        hours: float,
        staff_name: str | None = None,
        dry_run: bool = False,
    ) -> WriteResult:
        wb = openpyxl.load_workbook(self.path, keep_vba=self.path.suffix == ".xlsm")
        if self.sheet_name not in wb.sheetnames:
            return WriteResult(False, f"Sheet '{self.sheet_name}' not found in {self.path.name}")
        ws = wb[self.sheet_name]

        row, error = self._find_target_row(ws, entry_date, staff_name)
        if error:
            return WriteResult(False, error)

        if dry_run:
            return WriteResult(True, f"[dry run] would write row {row}: {code} / {task} / {hours}h", row)

        ws[f"{self.code_col}{row}"] = code
        ws[f"{self.task_col}{row}"] = task
        ws[f"{self.hours_col}{row}"] = hours

        self._save_with_retry(wb)
        return WriteResult(True, f"Wrote row {row}: {code} / {task} / {hours}h", row)

    def _save_with_retry(self, wb, attempts: int = 4, base_delay: float = 2.0):
        for attempt in range(attempts):
            try:
                wb.save(self.path)
                return
            except PermissionError:
                if attempt == attempts - 1:
                    raise PermissionError(
                        f"Could not save {self.path} — is it open in Excel? Close it and re-run."
                    )
                time.sleep(base_delay * (2 ** attempt))
